"""
报表生成模块
------------
支持导出处理结果为多种格式：
- CSV: 便于数据处理
- Excel (xlsx): 带格式的详细报表
- JSON: 结构化数据导出
- HTML: 可视化报表，可在浏览器查看
"""

import os
import json
import csv
import logging
from datetime import datetime
from typing import List, Optional
from pathlib import Path

from .processor import ProcessReport, BatchReport

logger = logging.getLogger("VideoDedup.Report")


class ReportExporter:
    """报表导出器"""

    def __init__(self, output_dir: str = ""):
        self.output_dir = output_dir or os.getcwd()

    def export_csv(self, batch_report: BatchReport, filepath: Optional[str] = None) -> str:
        """
        导出 CSV 报表。

        CSV 包含每个处理文件的详细结果。
        """
        if filepath is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.join(self.output_dir, f"video_dedup_report_{timestamp}.csv")

        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            # 表头
            writer.writerow([
                "文件名", "状态", "视频时长(秒)", "找到片段数", "匹配素材",
                "重复总时长(秒)", "重复占比(%)", "处理耗时(秒)", "错误信息"
            ])

            for report in batch_report.file_reports:
                writer.writerow([
                    report.video_name,
                    report.status,
                    f"{report.video_duration:.2f}",
                    report.segments_found,
                    "、".join(os.path.basename(p) for p in report.matched_references),
                    f"{report.total_match_duration:.2f}",
                    f"{(report.total_match_duration / report.video_duration * 100) if report.video_duration > 0 else 0:.1f}",
                    f"{report.processing_time:.2f}",
                    report.error_message,
                ])

            # 汇总行
            total_dur = sum(r.video_duration for r in batch_report.file_reports)
            writer.writerow([
                "【汇总】",
                f"处理 {batch_report.processed_files}/{batch_report.total_files} 文件",
                f"{total_dur:.2f}",
                batch_report.total_segments,
                "",
                f"{batch_report.total_match_duration:.2f}",
                f"{(batch_report.total_match_duration / total_dur * 100) if total_dur > 0 else 0:.1f}",
                f"{batch_report.total_processing_time:.2f}",
                "",
            ])

        logger.info(f"CSV 报表已导出: {filepath}")
        return filepath

    def export_json(self, batch_report: BatchReport, filepath: Optional[str] = None) -> str:
        """导出 JSON 报表（包含完整检测结果）"""
        if filepath is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.join(self.output_dir, f"video_dedup_report_{timestamp}.json")

        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

        data = {
            "export_time": datetime.now().isoformat(),
            "reference_video": batch_report.reference_path,
            "summary": {
                "total_files": batch_report.total_files,
                "processed_files": batch_report.processed_files,
                "files_with_matches": batch_report.files_with_matches,
                "total_segments": batch_report.total_segments,
                "total_match_duration": batch_report.total_match_duration,
                "total_processing_time": batch_report.total_processing_time,
            },
            "files": []
        }

        for report in batch_report.file_reports:
            file_data = {
                "path": report.video_path,
                "name": report.video_name,
                "duration": report.video_duration,
                "status": report.status,
                "segments_found": report.segments_found,
                "total_match_duration": report.total_match_duration,
                "processing_time": report.processing_time,
                "error": report.error_message,
                "matched_references": report.matched_references,
                "segments": [],
            }

            if report.detection:
                for seg in report.detection.segments:
                    file_data["segments"].append({
                        "ref_start": seg.ref_start_time,
                        "ref_end": seg.ref_end_time,
                        "target_start": seg.target_start_time,
                        "target_end": seg.target_end_time,
                        "similarity_avg": round(seg.avg_similarity, 1),
                        "similarity_min": round(seg.min_similarity, 1),
                        "duration": seg.duration,
                    })

            if report.cut_result:
                file_data["cut_output"] = report.cut_result.output_path
                file_data["cut_success"] = report.cut_result.success
                file_data["removed_duration"] = report.cut_result.removed_duration

            data["files"].append(file_data)

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

        logger.info(f"JSON 报表已导出: {filepath}")
        return filepath

    def export_excel(self, batch_report: BatchReport, filepath: Optional[str] = None) -> str:
        """
        导出 Excel 报表（xlsx 格式）。

        包含：
        - 汇总 sheet
        - 文件详情 sheet
        - 匹配片段详情 sheet
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.warning("openpyxl 未安装，无法导出 Excel 报表")
            return ""

        if filepath is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.join(self.output_dir, f"video_dedup_report_{timestamp}.xlsx")

        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

        wb = Workbook()

        # 样式
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thin", color="D0D0D0"),
        )

        # ===== Sheet 1: 汇总 =====
        ws_summary = wb.active
        ws_summary.title = "处理汇总"

        summary_data = [
            ["处理时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["参照视频", batch_report.reference_path],
            ["处理文件数", f"{batch_report.processed_files}/{batch_report.total_files}"],
            ["匹配文件数", batch_report.files_with_matches],
            ["发现重复片段总数", batch_report.total_segments],
            ["重复总时长(秒)", f"{batch_report.total_match_duration:.2f}"],
            ["总处理耗时(秒)", f"{batch_report.total_processing_time:.2f}"],
        ]
        for row_idx, (label, value) in enumerate(summary_data, 1):
            ws_summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2, value=str(value))

        ws_summary.column_dimensions['A'].width = 20
        ws_summary.column_dimensions['B'].width = 60

        # ===== Sheet 2: 文件详情 =====
        ws_files = wb.create_sheet("文件详情")

        file_headers = [
            "文件名", "状态", "时长(秒)", "片段数", "匹配素材",
            "重复时长(秒)", "重复占比", "处理耗时(秒)", "错误信息"
        ]
        for col, header in enumerate(file_headers, 1):
            cell = ws_files.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, report in enumerate(batch_report.file_reports, 2):
            values = [
                report.video_name,
                report.status,
                round(report.video_duration, 2),
                report.segments_found,
                "、".join(os.path.basename(p) for p in report.matched_references),
                round(report.total_match_duration, 2),
                f"{(report.total_match_duration / report.video_duration * 100):.1f}%" if report.video_duration > 0 else "0%",
                f"{report.processing_time:.1f}",
                report.error_message,
            ]
            for col, val in enumerate(values, 1):
                cell = ws_files.cell(row=row_idx, column=col, value=val)
                cell.alignment = cell_alignment
                cell.border = thin_border

        # 列宽
        widths = [40, 12, 12, 10, 30, 14, 12, 14, 30]
        for i, w in enumerate(widths, 1):
            ws_files.column_dimensions[get_column_letter(i)].width = w

        # ===== Sheet 3: 匹配片段详情 =====
        ws_segments = wb.create_sheet("匹配片段")

        seg_headers = [
            "文件名", "参照起始(s)", "参照结束(s)", "目标起始(s)",
            "目标结束(s)", "片段时长(s)", "平均相似度", "最低相似度"
        ]
        for col, header in enumerate(seg_headers, 1):
            cell = ws_segments.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        seg_row = 2
        for report in batch_report.file_reports:
            if report.detection:
                for seg in report.detection.segments:
                    values = [
                        report.video_name,
                        round(seg.ref_start_time, 2),
                        round(seg.ref_end_time, 2),
                        round(seg.target_start_time, 2),
                        round(seg.target_end_time, 2),
                        round(seg.duration, 2),
                        f"{seg.avg_similarity:.1f}%",
                        f"{seg.min_similarity:.1f}%",
                    ]
                    for col, val in enumerate(values, 1):
                        cell = ws_segments.cell(row=seg_row, column=col, value=val)
                        cell.alignment = cell_alignment
                        cell.border = thin_border
                    seg_row += 1

        seg_widths = [40, 14, 14, 14, 14, 14, 14, 14]
        for i, w in enumerate(seg_widths, 1):
            ws_segments.column_dimensions[get_column_letter(i)].width = w

        wb.save(filepath)
        logger.info(f"Excel 报表已导出: {filepath}")
        return filepath

    def export_html(self, batch_report: BatchReport, filepath: Optional[str] = None) -> str:
        """导出 HTML 可视化报表"""
        if filepath is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.join(self.output_dir, f"video_dedup_report_{timestamp}.html")

        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

        # 构建文件行
        file_rows = ""
        for report in batch_report.file_reports:
            status_class = {
                "completed": "status-ok",
                "no_match": "status-info",
                "error": "status-error",
            }.get(report.status, "status-warn")

            matched_refs = "、".join(os.path.basename(p) for p in report.matched_references) or "—"
            file_rows += f"""
            <tr>
                <td>{report.video_name}</td>
                <td><span class="badge {status_class}">{report.status}</span></td>
                <td class="num">{report.video_duration:.1f}s</td>
                <td class="num">{report.segments_found}</td>
                <td>{matched_refs}</td>
                <td class="num">{report.total_match_duration:.1f}s</td>
                <td class="num">{report.processing_time:.1f}s</td>
            </tr>"""

        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>视频去重处理报告</title>
<style>
    * {{ margin: 0; padding: 0; box-sizing: border-box; }}
    body {{ font-family: "Segoe UI", "Microsoft YaHei", sans-serif; background: #f5f5f5; color: #333; padding: 20px; }}
    .container {{ max-width: 1000px; margin: 0 auto; }}
    .header {{ background: linear-gradient(135deg, #0078d4, #005a9e); color: white; padding: 30px; border-radius: 12px; margin-bottom: 20px; }}
    .header h1 {{ font-size: 24px; margin-bottom: 8px; }}
    .header p {{ opacity: 0.85; font-size: 14px; }}
    .card {{ background: white; border-radius: 10px; padding: 24px; margin-bottom: 20px; box-shadow: 0 2px 8px rgba(0,0,0,0.08); }}
    .card h2 {{ font-size: 18px; margin-bottom: 16px; color: #0078d4; }}
    .summary-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 16px; }}
    .summary-item {{ background: #f8f9fa; padding: 16px; border-radius: 8px; text-align: center; }}
    .summary-item .value {{ font-size: 28px; font-weight: 700; color: #0078d4; }}
    .summary-item .label {{ font-size: 12px; color: #888; margin-top: 4px; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th {{ background: #f0f4f8; padding: 12px; text-align: left; font-weight: 600; border-bottom: 2px solid #e0e0e0; font-size: 13px; }}
    td {{ padding: 10px 12px; border-bottom: 1px solid #f0f0f0; font-size: 13px; }}
    tr:hover {{ background: #f8f9fa; }}
    .num {{ text-align: right; font-variant-numeric: tabular-nums; }}
    .badge {{ display: inline-block; padding: 3px 10px; border-radius: 12px; font-size: 12px; font-weight: 600; }}
    .status-ok {{ background: #d4edda; color: #155724; }}
    .status-info {{ background: #d1ecf1; color: #0c5460; }}
    .status-error {{ background: #f8d7da; color: #721c24; }}
    .status-warn {{ background: #fff3cd; color: #856404; }}
    .footer {{ text-align: center; color: #999; font-size: 12px; margin-top: 30px; }}
</style>
</head>
<body>
<div class="container">
    <div class="header">
        <h1>📹 视频去重处理报告</h1>
        <p>生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | 参照视频: {os.path.basename(batch_report.reference_path)}</p>
    </div>

    <div class="card">
        <h2>📊 处理汇总</h2>
        <div class="summary-grid">
            <div class="summary-item"><div class="value">{batch_report.processed_files}</div><div class="label">处理文件数</div></div>
            <div class="summary-item"><div class="value">{batch_report.files_with_matches}</div><div class="label">发现重复的文件</div></div>
            <div class="summary-item"><div class="value">{batch_report.total_segments}</div><div class="label">重复片段总数</div></div>
            <div class="summary-item"><div class="value">{batch_report.total_match_duration:.1f}s</div><div class="label">重复总时长</div></div>
            <div class="summary-item"><div class="value">{batch_report.total_processing_time:.1f}s</div><div class="label">总处理耗时</div></div>
        </div>
    </div>

    <div class="card">
        <h2>📁 文件处理详情</h2>
        <table>
            <thead><tr><th>文件名</th><th>状态</th><th>时长</th><th>片段</th><th>匹配素材</th><th>重复时长</th><th>耗时</th></tr></thead>
            <tbody>{file_rows}</tbody>
        </table>
    </div>

    <div class="footer">Video Dedup Tool v1.0 · 自动生成</div>
</div>
</body>
</html>"""

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html)

        logger.info(f"HTML 报表已导出: {filepath}")
        return filepath
